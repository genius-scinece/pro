#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> Word(.docx) 変換（見出し・表・箇条書き・コードブロック・引用に対応）"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

JP_FONT = "Yu Gothic"  # 游ゴシック（Wordで自然に表示。無ければ環境の既定に置換される）

def set_font(run, size=None, bold=None, color=None, mono=False):
    run.font.name = "Consolas" if mono else JP_FONT
    # 日本語グリフ用に東アジアフォントも指定
    rpr = run._element.get_or_add_rPr()
    from docx.oxml.ns import qn
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), "Consolas" if mono else JP_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)

def add_inline(paragraph, text):
    """**太字** と `コード` を解釈してrunを追加"""
    # トークン分割：**bold** と `code`
    parts = re.split(r'(\*\*.+?\*\*|`.+?`)', text)
    for p in parts:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            r = paragraph.add_run(p[2:-2]); set_font(r, bold=True)
        elif p.startswith('`') and p.endswith('`'):
            r = paragraph.add_run(p[1:-1]); set_font(r, mono=True, color=(0xC7, 0x25, 0x4E))
        else:
            r = paragraph.add_run(p); set_font(r)

def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    # 既定スタイルのフォント
    style = doc.styles['Normal']
    style.font.name = JP_FONT
    style.font.size = Pt(10.5)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # コードブロック
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i]); i += 1
            i += 1  # 閉じる```
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run('\n'.join(code_lines))
            set_font(r, size=9, mono=True)
            # 薄い背景（シェーディング）
            from docx.oxml.ns import qn
            pPr = p._p.get_or_add_pPr()
            shd = pPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'F2F2F2'})
            pPr.append(shd)
            continue

        # 表
        if line.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
            header = [c.strip() for c in line.strip().strip('|').split('|')]
            i += 2  # ヘッダ行と区切り行
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Light Grid Accent 1'
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].clear()
                add_inline(cell.paragraphs[0], h)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ''
                    cells[j].paragraphs[0].clear()
                    add_inline(cells[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        # 見出し
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            h = doc.add_heading('', level=min(level, 4))
            add_inline(h, text)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
            i += 1
            continue

        # 水平線
        if re.match(r'^\s*---+\s*$', line):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            from docx.oxml.ns import qn
            pbdr = pPr.makeelement(qn('w:pBdr'), {})
            bottom = pbdr.makeelement(qn('w:bottom'), {
                qn('w:val'): 'single', qn('w:sz'): '6',
                qn('w:space'): '1', qn('w:color'): 'CCCCCC'})
            pbdr.append(bottom); pPr.append(pbdr)
            i += 1
            continue

        # 引用
        if line.strip().startswith('>'):
            quote_lines = []
            while i < n and lines[i].strip().startswith('>'):
                quote_lines.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            text = ' '.join(l for l in quote_lines if l.strip())
            if text.strip():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                add_inline(p, text)
                for run in p.runs:
                    run.font.italic = True
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # 箇条書き
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            indent = len(m.group(1))
            p = doc.add_paragraph(style='List Bullet')
            if indent >= 2:
                p.paragraph_format.left_indent = Pt(36)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 番号付き
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_inline(p, m.group(2))
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 通常段落
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print(f"保存: {docx_path}")

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
